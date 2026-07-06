import openpyxl
from django.contrib.auth.decorators import login_required
from django.core.cache import cache
from django.http import (
    HttpResponse,
    JsonResponse,
)
from django.shortcuts import render
from django.utils.html import strip_tags
from hitcount.utils import get_hitcount_model

from yummyrecipes.tasks.weather_tasks import fetch_weather_and_cache

from ..forms import (
    LocationForm,
)
from ..models import (
    BlogHistory,
    CreditHistory,
    CreditSpentHistory,
    Feedback,
    Follow,
    RedeemedCredit,
    UserLocation,
    UserProfile,
    comments,
    post,
    profile,
)


def email_template(request):
    return render(request, "yummyrecipes/email_template.html")


def handling_404(request, exception):
    return render(request, "yummyrecipes/404.html")


def save_location(request):
    if request.method == "POST":
        form = LocationForm(request.POST)

        if form.is_valid():
            location = form.cleaned_data["location"].strip().lower()

            # Update user profile
            user_profile, created = UserLocation.objects.get_or_create(user=request.user)
            user_profile.location = location
            user_profile.save()

            weather_cache_key = f"weather_data_{location}"
            lock_key = f"weather_fetching_{location}"
            recipe_cache_key = f"suggested_recipes_{location}"

            # Check if weather and recipe cache exist
            cached_weather_data = cache.get(weather_cache_key)
            cached_recipes = cache.get(recipe_cache_key)

            if not cached_weather_data and not cache.get(lock_key):
                cache.set(lock_key, True, timeout=30)
                fetch_weather_and_cache.delay(location)

            if not cached_weather_data:
                status = "fetching"
            elif not cached_recipes and not cached_recipes:
                status = "partial"
            else:
                status = "ready"

            suggested_recipes = []
            gemini_response = ""

            if cached_recipes:
                suggested_recipes = cached_recipes[0]
                gemini_response = cached_recipes[1]

            return JsonResponse(
                {
                    "status": status,
                    "weather_data": cached_weather_data,
                    "suggested_recipes": suggested_recipes,
                    "gemini_response": gemini_response,
                }
            )

    else:
        form = LocationForm()

    return render(request, "yummyrecipes/home.html", {"form": form})


@login_required
def export_user_data(request):
    if request.user.is_authenticated:
        user = request.user
        full_name = f"{user.first_name} {user.last_name}".strip() or "user_data"
        safe_filename = full_name.replace(" ", "_")
        filename = f"{safe_filename}_data.xlsx"
        workbook = openpyxl.Workbook()

        # Sheet for recipes
        sheet_recipes = workbook.active
        sheet_recipes.title = "Recipes"
        sheet_recipes.append(
            [
                "Title",
                "Timing (in mins)",
                "Servings",
                "Type",
                "Cuisine",
                "Category",
                "Difficulty",
                "Ingredients",
                "Instructions",
                "Date Posted",
                "Date Modified",
                "Likes",
                "Dislikes",
                "Views",
            ]
        )

        for data in post.objects.filter(author=user):
            clean_content = strip_tags(data.content)
            clean_ingredients = strip_tags(data.ingredients)
            hit_count = get_hitcount_model().objects.get_for_object(data).hits

            sheet_recipes.append(
                [
                    data.title,
                    data.timing,
                    data.servings,
                    data.type,
                    data.cuisine,
                    data.category,
                    data.difficulty,
                    clean_ingredients,
                    clean_content,
                    data.date_post.strftime("%d-%m-%Y %I:%M %p"),
                    data.date_modified.strftime("%d-%m-%Y %I:%M %p"),
                    data.likes.count(),
                    data.dislikes.count(),
                    hit_count,
                ]
            )

        # Sheet for comments received on user’s recipes
        sheet_comments = workbook.create_sheet(title="Comments")
        sheet_comments.append(["Recipe Title", "Comment", "Commented By", "Date"])

        for comment in comments.objects.filter(post_super__author=user):
            sheet_comments.append(
                [
                    comment.post_super.title,
                    comment.comment,
                    f"{comment.comment_user.first_name} {comment.comment_user.last_name}".strip(),
                    comment.date_comment.strftime("%d-%m-%Y %I:%M %p"),
                ]
            )

        # Sheet for recipes visited (history)
        sheet_history = workbook.create_sheet(title="Recipes Visited")
        sheet_history.append(["Recipe Title", "Timestamp"])

        for history in BlogHistory.objects.filter(user=user):
            sheet_history.append(
                [
                    history.blog_post.title,
                    history.timestamp.strftime("%d-%m-%Y %I:%M %p"),
                ]
            )

        # Sheet for redeemed credits
        sheet_redeemed = workbook.create_sheet(title="Redeemed Credits")
        sheet_redeemed.append(["Amount", "Redeemed Timestamp"])

        for credit in RedeemedCredit.objects.filter(user=user):
            sheet_redeemed.append([credit.amount, credit.redeemed_timestamp.strftime("%d-%m-%Y %I:%M %p")])

        # Sheet for credit history
        sheet_credit_history = workbook.create_sheet(title="Credit History")
        sheet_credit_history.append(["Amount", "Action", "Timestamp"])

        for history in CreditHistory.objects.filter(user=user):
            sheet_credit_history.append(
                [
                    history.amount,
                    history.get_credit_action_display(),
                    history.earned_timestamp.strftime("%d-%m-%Y %I:%M %p"),
                ]
            )

        # Sheet for credit spent history
        sheet_spent_history = workbook.create_sheet(title="Credit Spent")
        sheet_spent_history.append(["Recipe Title", "Amount Spent", "Timestamp"])

        for spent in CreditSpentHistory.objects.filter(user=user):
            sheet_spent_history.append(
                [
                    spent.recipename.title,
                    spent.amount,
                    spent.spent_timestamp.strftime("%d-%m-%Y %I:%M %p"),
                ]
            )

        # Sheet for feedback
        sheet_feedback = workbook.create_sheet(title="Feedback")
        sheet_feedback.append(["Name", "Email", "Subject", "Message", "Timestamp"])

        for feedback in Feedback.objects.filter(user=user):
            sheet_feedback.append(
                [
                    feedback.name,
                    feedback.email,
                    feedback.subject,
                    feedback.message,
                    feedback.timestamp.strftime("%d-%m-%Y %I:%M %p"),
                ]
            )

        # Sheet for profile data
        sheet_profile_theme = workbook.create_sheet(title="Profile data")
        sheet_profile_theme.append(
            [
                "Full Name",
                "Email",
                "Last Login",
                "Joined Date",
                "Bio",
                "Date of Birth",
                "Gender",
                "Credits",
                "Earned Credits",
                "Redeemed Credits",
                "Credits Spent",
                "Facebook",
                "Instagram",
                "Twitter",
                "Threads",
                "YouTube",
                "Website",
                "Preference Type",
                "Preference Category",
                "Preference Cuisine",
                "",
                "",
                "Primary Color",
                "Secondary Color",
                "Tertiary Color",
                "Active Link Color",
                "Hover Color",
                "Neutral Primary Color",
                "Neutral Secondary Color",
            ]
        )

        user_profile = profile.objects.filter(user=user).first()
        user_theme = UserProfile.objects.filter(user=user).first()
        last_login = user.last_login.strftime("%d-%m-%Y %I:%M %p") if user.last_login else "N/A"
        date_joined = user.date_joined.strftime("%d-%m-%Y %I:%M %p")

        if user_profile and user_theme:
            sheet_profile_theme.append(
                [
                    f"{user.first_name} {user.last_name}".strip(),
                    user.email,
                    last_login,
                    date_joined,
                    user_profile.bio,
                    user_profile.dob.strftime("%d-%m-%Y") if user_profile.dob else "N/A",
                    user_profile.gender,
                    user_profile.credits,
                    user_profile.earned_credits,
                    user_profile.redeemed_credits,
                    user_profile.credits_spent,
                    user_profile.facebook,
                    user_profile.instagram,
                    user_profile.twitter,
                    user_profile.threads,
                    user_profile.youtube,
                    user_profile.website,
                    user_profile.preference_type if user_profile.preference_type != "Select" else "N/A",
                    user_profile.preference_category if user_profile.preference_category != "Select" else "N/A",
                    user_profile.preference_cuisine if user_profile.preference_cuisine != "Select" else "N/A",
                    "",
                    "",
                    user_theme.primary_color,
                    user_theme.secondary_color,
                    user_theme.tertiary_color,
                    user_theme.active_link_color,
                    user_theme.hover_color,
                    user_theme.neutral_primary,
                    user_theme.neutral_secondary,
                ]
            )

        # Sheet for Follows (Followers & Following)
        sheet_follows = workbook.create_sheet(title="Follows")
        sheet_follows.append(["Follower", "Following", "Timestamp"])

        for follow in Follow.objects.filter(follower=user):
            follower_name = f"{follow.follower.first_name} {follow.follower.last_name}".strip()
            following_name = f"{follow.following.first_name} {follow.following.last_name}".strip()

            if follow.follower == user:
                follower_name += " (You)"
            if follow.following == user:
                following_name += " (You)"

            sheet_follows.append(
                [
                    follower_name,
                    following_name,
                    follow.created_at.strftime("%d-%m-%Y %I:%M %p"),
                ]
            )

        for follow in Follow.objects.filter(following=user):
            follower_name = f"{follow.follower.first_name} {follow.follower.last_name}".strip()
            following_name = f"{follow.following.first_name} {follow.following.last_name}".strip()

            if follow.follower == user:
                follower_name += " (You)"
            if follow.following == user:
                following_name += " (You)"

            sheet_follows.append(
                [
                    follower_name,
                    following_name,
                    follow.created_at.strftime("%d-%m-%Y %I:%M %p"),
                ]
            )

        # **Prepare Response**
        response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'

        workbook.save(response)
        return response
    else:
        return HttpResponse("Unauthorized", status=401)
